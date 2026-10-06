# -*- coding: utf-8 -*-
"""分析两组全量(20ms)实验 CSV, 计算"极速→速度变负"耗时, 并生成 Excel 速度-时间折线图。

输入: speed_log_group1_no_neutral.csv / speed_log_group2_neutral.csv
输出: WTBlackList_gear_analysis.xlsx + 控制台分析
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
G1 = os.path.join(BASE, "speed_log_group1_no_neutral.csv")
G2 = os.path.join(BASE, "speed_log_group2_neutral.csv")
OUT = os.path.join(BASE, "WTBlackList_gear_analysis.xlsx")

FILL = PatternFill("solid", fgColor="2F5496")
FH = Font(color="FFFFFF", bold=True)


def load(path):
    """新 CSV 列: time_s, rtt_ms, speed, gear, gear_neutral, rpm, vario, altitude"""
    rows = []
    with open(path, encoding="utf-8-sig") as f:
        r = csv.reader(f)
        next(r, None)  # 表头
        for line in r:
            try:
                t = float(line[0]); s = float(line[2])
                g = line[3] if len(line) > 3 and line[3] != "" else None
                rp = float(line[5]) if len(line) > 5 and line[5] != "" else None
                rows.append((t, s, g, rp))
            except (ValueError, IndexError):
                continue
    return rows


def analyze(rows, label):
    """极速段结束(最后一个峰值点) → 速度变负时刻 → 耗时。"""
    vals = [s for _t, s, _g, _r in rows if s is not None]
    if not vals:
        return None
    peak_val = max(vals)
    if peak_val < 40:
        return None
    # 极速段结束 = 最后一个达到峰值的采样点(之后速度开始下降 = 用户开始换向)
    peak = None
    for t, s, _g, _r in rows:
        if s is not None and s >= peak_val:
            peak = (t, s)
    neg = None
    for t, s, _g, _r in rows:
        if t > peak[0] and s is not None and s < 0:
            neg = (t, s)
            break
    # 切入倒挡: 极速后第一个挡位 <= 3 (gear_neutral=4, gear=3 即 UI -1 倒挡)
    rev = None
    for t, _s, g, _r in rows:
        if t > peak[0] and g is not None:
            try:
                if float(g) <= 3:
                    rev = (t, g)
                    break
            except ValueError:
                continue
    if neg is None:
        return None
    return {
        "label": label,
        "peak_t": peak[0], "peak_s": peak[1],
        "neg_t": neg[0], "neg_s": neg[1],
        "dur": round(neg[0] - peak[0], 2),
        "rev_t": rev[0] if rev else None, "rev_g": rev[1] if rev else None,
        "rev_dur": round(rev[0] - peak[0], 2) if rev else None,
    }


def write_chart_sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    ws["A1"] = "时间(秒)"
    ws["B1"] = "速度(km/h)"
    ws["C1"] = "挡位(8111)"
    for c in ("A1", "B1", "C1"):
        ws[c].fill = FILL
        ws[c].font = FH
    for i, (t, s, g, rp) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=t)
        ws.cell(row=i, column=2, value=s)
        ws.cell(row=i, column=3, value=g)
    n = len(rows)
    sc = ScatterChart()
    sc.title = f"{name} 速度-时间(20ms全量)"
    sc.style = 13
    sc.x_axis.title = "时间(秒)"
    sc.y_axis.title = "速度(km/h)"
    sc.y_axis.scaling.min = -40
    sc.y_axis.scaling.max = 80
    sc.height = 13
    sc.width = 26
    xv = Reference(ws, min_col=1, min_row=2, max_row=n + 1)
    yv = Reference(ws, min_col=2, min_row=2, max_row=n + 1)
    s = Series(yv, xv, title=f"{name} 速度")
    s.marker.symbol = "circle"
    s.marker.size = 3
    sc.series.append(s)
    ws.add_chart(sc, "E2")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10


def main() -> int:
    g1 = load(G1)
    g2 = load(G2)
    r1 = analyze(g1, "不用空挡")
    r2 = analyze(g2, "用空挡")

    print(f"第一组(不用空挡) 采样点: {len(g1)}")
    print(f"第二组(用空挡)   采样点: {len(g2)}\n")

    wb = Workbook()
    wb.remove(wb.active)

    for name, rows in (("第一组-不用空挡", g1), ("第二组-用空挡", g2)):
        write_chart_sheet(wb, name, rows)

    ws = wb.create_sheet("对比-极速到速度变负")
    hdr = ["组别", "极速时刻(s)", "极速(km/h)", "速度变负时刻(s)", "速度(km/h)", "耗时(s)",
           "切入倒挡时刻(s)", "挡位(8111)", "切入倒挡耗时(s)"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = FILL
        c.font = FH
    for i, r in enumerate((r1, r2), start=2):
        if r is None:
            continue
        ws.cell(row=i, column=1, value=r["label"])
        ws.cell(row=i, column=2, value=r["peak_t"])
        ws.cell(row=i, column=3, value=r["peak_s"])
        ws.cell(row=i, column=4, value=r["neg_t"])
        ws.cell(row=i, column=5, value=r["neg_s"])
        ws.cell(row=i, column=6, value=r["dur"])
        ws.cell(row=i, column=7, value=r.get("rev_t"))
        ws.cell(row=i, column=8, value=r.get("rev_g"))
        ws.cell(row=i, column=9, value=r.get("rev_dur"))
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    wb.save(OUT)

    if r1:
        print(f"【不用空挡】极速 {r1['peak_s']}km/h @{r1['peak_t']:.1f}s → "
              f"速度变负 {r1['neg_s']}km/h @{r1['neg_t']:.1f}s = {r1['dur']}s | "
              f"切入倒挡(挡位≤3) @{r1.get('rev_t')}s = {r1.get('rev_dur')}s")
    if r2:
        print(f"【用空挡】  极速 {r2['peak_s']}km/h @{r2['peak_t']:.1f}s → "
              f"速度变负 {r2['neg_s']}km/h @{r2['neg_t']:.1f}s = {r2['dur']}s | "
              f"切入倒挡(挡位≤3) @{r2.get('rev_t')}s = {r2.get('rev_dur')}s")
    if r1 and r2:
        diff = round(r1["dur"] - r2["dur"], 2)
        print(f"\n差异(速度变负): 用空挡比不用空挡快 {diff}s ({(1 - r2['dur']/r1['dur'])*100:.0f}%)")
        if r1.get("rev_dur") and r2.get("rev_dur"):
            rd = round(r1["rev_dur"] - r2["rev_dur"], 2)
            print(f"差异(切入倒挡): 用空挡比不用空挡快 {rd}s "
                  f"({(1 - r2['rev_dur']/r1['rev_dur'])*100:.0f}%)")
    print("\n已生成 Excel:", OUT)
    return 0


if __name__ == "__main__":
    main()
