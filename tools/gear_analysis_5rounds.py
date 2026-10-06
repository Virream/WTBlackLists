# -*- coding: utf-8 -*-
"""5 轮重复实验聚合分析: 豹2PL 换挡对比(不用空挡 vs 用空挡)。

每组 5 轮 CSV (speed_log_g1_run1~5 / speed_log_g2_run1~5),
对每轮计算"极速→速度变负耗时"和"切入倒挡耗时"(挡位≤3),
再对每组求 平均±标准差, 输出对比与 Excel。
"""
import csv
import os
import statistics

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

G1_FILES = [f"speed_log_g1_run{i}.csv" for i in range(1, 6)]
G2_FILES = [f"speed_log_g2_run{i}.csv" for i in range(1, 6)]
OUT = os.path.join(BASE, "WTBlackList_gear_5rounds.xlsx")

FILL = PatternFill("solid", fgColor="2F5496")
FH = Font(color="FFFFFF", bold=True)


def load(path):
    """新 CSV 列: time_s, rtt_ms, speed, gear, gear_neutral, rpm, vario, altitude"""
    rows = []
    with open(path, encoding="utf-8-sig") as f:
        r = csv.reader(f)
        next(r, None)  # 表头
        for line in r:
            try:
                t = float(line[0]); s = float(line[2])
                g = line[3] if len(line) > 3 and line[3] != "" else None
                rows.append((t, s, g))
            except (ValueError, IndexError):
                continue
    return rows


def analyze(rows):
    """极速段结束(最后一个峰值点) → 速度变负时刻 → 耗时; 切入倒挡(挡位≤3)耗时。"""
    vals = [s for _t, s, _g in rows if s is not None]
    if not vals:
        return None
    peak_val = max(vals)
    if peak_val < 40:
        return None
    # 极速段结束 = 最后一个达到峰值的采样点(之后速度开始下降 = 用户开始换向)
    peak = None
    for t, s, _g in rows:
        if s is not None and s >= peak_val:
            peak = (t, s)
    neg = None
    for t, s, _g in rows:
        if t > peak[0] and s is not None and s < 0:
            neg = (t, s)
            break
    rev = None
    for t, _s, g in rows:
        if t > peak[0] and g is not None:
            try:
                if float(g) <= 3:
                    rev = (t, g)
                    break
            except ValueError:
                continue
    if neg is None:
        return None
    return {
        "peak_t": peak[0], "peak_s": peak[1],
        "neg_t": neg[0], "neg_s": neg[1],
        "dur": neg[0] - peak[0],
        "rev_t": rev[0] if rev else None, "rev_g": rev[1] if rev else None,
        "rev_dur": rev[0] - peak[0] if rev else None,
    }


def agg(results):
    """对一组多个结果求 dur / rev_dur 的均值与标准差。"""
    durs = [r["dur"] for r in results if r is not None]
    revs = [r["rev_dur"] for r in results if r is not None and r.get("rev_dur") is not None]
    return {
        "n": len(durs),
        "dur_mean": statistics.mean(durs) if durs else None,
        "dur_std": statistics.stdev(durs) if len(durs) > 1 else 0.0,
        "rev_mean": statistics.mean(revs) if revs else None,
        "rev_std": statistics.stdev(revs) if len(revs) > 1 else 0.0,
    }


def fmt(v, nd=2):
    return f"{v:.{nd}f}" if v is not None else "-"


def main() -> int:
    groups = [
        ("不用空挡", G1_FILES, "g1"),
        ("用空挡", G2_FILES, "g2"),
    ]
    all_rows = {}   # (g, i) -> results
    per_group = {}

    for label, files, tag in groups:
        results = []
        for i, fn in enumerate(files, start=1):
            path = os.path.join(BASE, fn)
            rows = load(path)
            r = analyze(rows)
            results.append(r)
            all_rows[(tag, i)] = (r, rows)
            print(f"[{label}] run{i}: {('OK' if r else '无效(未变负/极速不足)')} "
                  f"样本={len(rows)}")
        per_group[label] = agg(results)

    print("\n=== 每组聚合(平均±标准差) ===")
    for label in ("不用空挡", "用空挡"):
        a = per_group[label]
        print(f"{label}: 有效{a['n']}轮 | 速度变负耗时 {fmt(a['dur_mean'])}s ± {fmt(a['dur_std'])}s"
              f" | 切入倒挡耗时 {fmt(a['rev_mean'])}s ± {fmt(a['rev_std'])}s")
    d1 = per_group["不用空挡"]; d2 = per_group["用空挡"]
    if d1["dur_mean"] and d2["dur_mean"]:
        diff = d1["dur_mean"] - d2["dur_mean"]
        pct = (1 - d2["dur_mean"] / d1["dur_mean"]) * 100
        print(f"\n速度变负: 用空挡比不用空挡快 {fmt(diff)}s ({pct:.1f}%)")
    if d1["rev_mean"] and d2["rev_mean"]:
        rdiff = d1["rev_mean"] - d2["rev_mean"]
        rpct = (1 - d2["rev_mean"] / d1["rev_mean"]) * 100
        print(f"切入倒挡: 用空挡比不用空挡快 {fmt(rdiff)}s ({rpct:.1f}%)")

    # ===== Excel =====
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet1 每轮明细
    ws = wb.create_sheet("每轮明细")
    hdr = ["组别", "轮次", "极速时刻(s)", "极速(km/h)", "速度变负时刻(s)",
           "速度变负耗时(s)", "切入倒挡时刻(s)", "挡位(8111)", "切入倒挡耗时(s)", "状态"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = FILL; c.font = FH
    rowi = 2
    for label, _files, tag in groups:
        for i in range(1, 6):
            r, _rows = all_rows[(tag, i)]
            vals = [label, f"run{i}"]
            if r is None:
                vals += ["-", "-", "-", "-", "-", "-", "-", "无效"]
            else:
                vals += [round(r["peak_t"], 2), r["peak_s"], round(r["neg_t"], 2),
                         round(r["dur"], 2),
                         (round(r["rev_t"], 2) if r["rev_t"] else "-"),
                         (r["rev_g"] if r["rev_g"] else "-"),
                         (round(r["rev_dur"], 2) if r["rev_dur"] else "-"), "OK"]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=rowi, column=j, value=v)
            rowi += 1
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18

    # Sheet2 汇总对比
    ws2 = wb.create_sheet("汇总对比")
    hdr2 = ["组别", "有效轮数", "速度变负均值(s)", "速度变负标准差(s)",
            "切入倒挡均值(s)", "切入倒挡标准差(s)"]
    for j, h in enumerate(hdr2, start=1):
        c = ws2.cell(row=1, column=j, value=h)
        c.fill = FILL; c.font = FH
    for ri, label in enumerate(("不用空挡", "用空挡"), start=2):
        a = per_group[label]
        ws2.cell(row=ri, column=1, value=label)
        ws2.cell(row=ri, column=2, value=a["n"])
        ws2.cell(row=ri, column=3, value=a["dur_mean"])
        ws2.cell(row=ri, column=4, value=a["dur_std"])
        ws2.cell(row=ri, column=5, value=a["rev_mean"])
        ws2.cell(row=ri, column=6, value=a["rev_std"])
    # 柱状图: 速度变负耗时对比
    ch = BarChart()
    ch.type = "col"
    ch.style = 10
    ch.title = "极速→速度变负 耗时对比(平均±标准差)"
    ch.y_axis.title = "秒"
    data = Reference(ws2, min_col=3, min_row=1, max_row=3, max_col=4)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=3)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.height = 9; ch.width = 16
    ws2.add_chart(ch, "H2")
    for j in range(1, len(hdr2) + 1):
        ws2.column_dimensions[get_column_letter(j)].width = 22

    wb.save(OUT)
    print("\n已生成 Excel:", OUT)
    return 0


if __name__ == "__main__":
    main()
